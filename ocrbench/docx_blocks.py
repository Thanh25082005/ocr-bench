"""DOCX dạng bảng khối: mỗi trang là một bảng 2 cột — Loại (catalog) | Nội dung — mỗi dòng một khối model trả
về, theo thứ tự đọc từ trên xuống dưới.

Nội dung theo loại khối:
- Title / Section-header → chữ kiểu heading; Text, List-item, Caption, Footnote, Page-header, Page-footer → đoạn văn
- Table → bảng Word lồng trong ô (từ HTML của model); Formula → chuỗi LaTeX; Picture → ảnh cắt từ trang gốc theo bbox
Không cố dựng lại bố cục trang (việc đó là của ocrbench/docx_exact.py).
"""

from __future__ import annotations

import io
from pathlib import Path

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .docx_export import _add_markdown, _add_runs, _bidi_from_image, _set_bidi, new_document, para_rtl

_HEADING_LEVEL = {"Title": 1, "Section-header": 2}
_LABEL_COLOR = {"Title": (0xDC, 0x26, 0x26), "Section-header": (0x08, 0x91, 0xB2), "Table": (0x93, 0x33, 0xEA),
                "Picture": (0xD9, 0x77, 0x06), "Signature": (0xB4, 0x53, 0x09), "Formula": (0xDB, 0x27, 0x77)}
_DEFAULT_COLOR = (0x16, 0xA3, 0x4A)
_LABEL_W, _CONTENT_W = Cm(3.5), Cm(13.0)  # A4 dọc, lề mặc định ~16.5 cm bề rộng chữ


class _CellTarget:
    """Cho các hàm dựng nội dung của docx_export (vốn ghi vào Document) ghi vào một ô bảng."""

    def __init__(self, cell):
        self.cell = cell

    def add_paragraph(self, text: str = "", style=None):
        return self.cell.add_paragraph(text, style)

    def add_heading(self, text: str = "", level: int = 1):
        p = self.cell.add_paragraph(text, f"Heading {level}")
        p.paragraph_format.space_before = Pt(0)  # khoảng trống trên heading làm dòng bảng cao thừa
        return p

    def add_table(self, rows: int, cols: int):
        return self.cell.add_table(rows, cols)


def _shade(cell, hex_fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def _new_table(doc):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width, table.columns[1].width = _LABEL_W, _CONTENT_W  # w:gridCol — LibreOffice đọc độ rộng ở đây
    hdr = table.rows[0]
    trpr = hdr._tr.get_or_add_trPr()  # lặp dòng tiêu đề khi bảng sang trang mới
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)
    for cell, text in zip(hdr.cells, ("Loại (catalog)", "Nội dung")):
        cell.paragraphs[0].add_run(text).bold = True
        _shade(cell, "D9D9D9")
    return table


def _add_row(table, category: str):
    row = table.add_row()
    label, content = row.cells
    label.width, content.width = _LABEL_W, _CONTENT_W
    run = label.paragraphs[0].add_run(category)
    run.bold = True
    run.font.color.rgb = RGBColor(*_LABEL_COLOR.get(category, _DEFAULT_COLOR))
    return content


def _add_picture(cell, block: dict, image) -> bool:
    x1, y1, x2, y2 = [int(round(v)) for v in block["bbox"]]
    x1, y1, x2, y2 = max(0, x1), max(0, y1), min(image.width, x2), min(image.height, y2)
    if x2 - x1 < 3 or y2 - y1 < 3:
        return False
    buf = io.BytesIO()
    image.crop((x1, y1, x2, y2)).convert("RGB").save(buf, format="PNG")
    buf.seek(0)
    # cỡ thật theo tỉ lệ khổ trang A4 (21 cm), không vượt bề rộng cột nội dung
    width = min(int(Cm(21) * (x2 - x1) / image.width), int(_CONTENT_W - Cm(0.4)))
    cell.paragraphs[-1].add_run().add_picture(buf, width=width)
    return True


def _fill(cell, block: dict, image, cache: dict, blocks: list) -> None:
    category = block.get("category") or "Text"
    text = block.get("text") or ""
    if category in ("Picture", "Signature") or (not text.strip() and category != "Table"):
        bbox = block.get("bbox")
        if not (image is not None and bbox and len(bbox) == 4 and _add_picture(cell, block, image)):
            cell.paragraphs[0].add_run("(ảnh — không cắt được từ trang)")
        return
    if image is not None:
        text = _bidi_from_image(block, image, cache, blocks)
    target = _CellTarget(cell)
    if category in _HEADING_LEVEL:
        p = target.add_heading(level=_HEADING_LEVEL[category])
        _add_runs(p, text.lstrip("#").strip())
        if para_rtl(text):
            _set_bidi(p)
    elif category == "Formula":
        target.add_paragraph().add_run(text.strip()).font.name = "Consolas"
    elif not _add_markdown(target, text):
        cell.paragraphs[0].add_run("(trống)")
    _trim(cell)


def _trim(cell) -> None:
    """Ô mới luôn có sẵn một đoạn rỗng ở đầu; nội dung thêm sau nó → bỏ đoạn rỗng đó. Đoạn rỗng cuối (sau bảng
    lồng) cũng bỏ nếu phần tử cuối vẫn còn là đoạn (Word bắt buộc ô kết thúc bằng một đoạn)."""
    first = cell.paragraphs[0]
    if not first.text and not first._p.xpath(".//w:drawing") and len(cell._tc) > 2:
        first._p.getparent().remove(first._p)
    kids = [k for k in cell._tc if k.tag != qn("w:tcPr")]
    if len(kids) >= 3 and kids[-1].tag == qn("w:p") and kids[-2].tag == qn("w:p") \
            and not "".join(kids[-1].itertext()).strip():
        cell._tc.remove(kids[-1])


def add_page(doc, page, first: bool = False) -> None:
    """`page`: PageResult của Converter (dùng index, source, note, error, text, blocks, image)."""
    if not first:
        doc.add_page_break()
    doc.add_heading(f"Trang {page.index}", level=1)
    meta = f"nguồn: {page.source}" + (f" · LỖI: {page.error}" if page.error else "") \
        + (f" · ⚠ CẦN SOÁT: {page.note}" if page.note else "")
    run = doc.add_paragraph().add_run(meta)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    blocks = page.blocks or []
    if not blocks:  # lớp chữ PDF / chế độ chỉ chữ: không có khối → cả trang là một dòng Text
        if not page.text.strip():
            doc.add_paragraph("(Không đọc được chữ nào trên trang này)")
            return
        blocks = [{"category": "Text (cả trang)", "text": page.text}]
    table = _new_table(doc)
    cache: dict = {}
    for block in blocks:
        _fill(_add_row(table, block.get("category") or "Text"), block, page.image, cache, blocks)
    for row in table.rows:  # Word đọc độ rộng theo từng ô
        row.cells[0].width, row.cells[1].width = _LABEL_W, _CONTENT_W


def build_blocks_docx(pages: list, out_path: str | Path, title: str | None = None) -> Path:
    doc = new_document(title=title)
    for i, page in enumerate(pages):
        add_page(doc, page, first=(i == 0))
    out_path = Path(out_path)
    doc.save(out_path)
    return out_path
