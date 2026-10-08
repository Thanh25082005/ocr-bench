import yaml
from docx import Document

from tests.test_docx_exact import _page


def test_blocks_docx_is_category_content_table(tmp_path):
    """<tên>_khoi.docx: mỗi trang một bảng 2 cột Loại | Nội dung, mỗi dòng một khối theo thứ tự đọc;
    Table → bảng lồng trong ô, Picture → ảnh cắt từ trang."""
    (tmp_path / "m.jsonl").write_text("")
    (tmp_path / "c.yaml").write_text(yaml.safe_dump({"dataset": "m.jsonl", "models": [
        {"name": "lay", "adapter": "tests.test_docx_blocks:FakeWithPicture"}]}))
    from ocrbench.convert import Converter

    img, _ = _page()
    img.save(tmp_path / "a.png")
    res = Converter(tmp_path / "c.yaml", "lay").convert_file(tmp_path / "a.png", tmp_path / "o", force_ocr=True,
                                                            exact=False)
    assert res.docx_exact is None and res.docx_blocks.name == "a_khoi.docx"
    doc = Document(res.docx_blocks)
    assert doc.paragraphs[0].text == "Trang 1" and len(doc.tables) == 1
    rows = doc.tables[0].rows
    blocks = FakeWithPicture().predict(img, None).extra["blocks"]
    assert [r.cells[0].text for r in rows] == ["Loại (catalog)"] + [b["category"] for b in blocks]
    assert rows[1].cells[1].text == "INVOICE No. 2041"
    assert "Al Noor Trading LLC" in rows[2].cells[1].text
    nested = rows[4].cells[1].tables
    assert len(nested) == 1 and nested[0].cell(1, 2).text == "1,250.00"
    pics = [r for r in rows if r.cells[0].text == "Picture"]
    assert len(pics) == 2 and all(r.cells[1]._tc.xpath(".//w:drawing") for r in pics)


from ocrbench.adapters.base import Adapter, Prediction  # noqa: E402


class FakeWithPicture(Adapter):
    def predict(self, image, item):
        _, blocks = _page()
        blocks.append({"category": "Picture", "bbox": [10, 10, 80, 60]})
        return Prediction("x", extra={"layout": "ok", "blocks": blocks})
