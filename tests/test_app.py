"""PDF/ảnh → DOCX: phân tích kết quả OCR, ghi DOCX (bảng, chiều phải-sang-trái), chọn lớp chữ hay OCR."""

import random

import pytest
import yaml
from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageDraw

from ocrbench.adapters.base import Adapter, Prediction
from ocrbench.convert import Converter, parse_pages
from ocrbench.docx_export import add_page, is_rtl, new_document, parse_blocks

OCR_OUT = """# فاتورة ضريبية

**رقم الفاتورة:** INV-2041
العميل: شركة الأفق للتجارة

<table><tr><th colspan="2">البند والكمية</th><th>السعر</th></tr><tr><td>حاسوب</td><td>2</td><td>1,250.00</td></tr></table>

| Item | Qty |
|---|---|
| Pen | 3 |

Thank you for your business."""


class FakeOCR(Adapter):
    def predict(self, image, item):
        return Prediction(OCR_OUT + f"\n\n[{item.gt_type}]")


def test_parse_blocks():
    blocks = parse_blocks(OCR_OUT)
    kinds = [b.kind for b in blocks]
    assert kinds == ["heading", "para", "table", "table", "para"]
    assert blocks[2].rows[0] == [("البند والكمية", 2), ("السعر", 1)]
    assert blocks[3].rows == [[("Item", 1), ("Qty", 1)], [("Pen", 1), ("3", 1)]]
    assert is_rtl(blocks[1].text) and not is_rtl(blocks[4].text)


def test_docx_tables_rtl_and_pages(tmp_path):
    doc = new_document("t")
    add_page(doc, OCR_OUT, header="Trang 1", first=True)
    add_page(doc, "Page two in English.", header="Trang 2")
    doc.save(tmp_path / "x.docx")
    d = Document(tmp_path / "x.docx")
    assert len(d.tables) == 2
    t0 = d.tables[0]
    assert t0._tbl.tblPr.find(qn("w:bidiVisual")) is not None  # bảng tiếng Ả Rập: cột phải → trái
    assert t0.cell(0, 0).text == "البند والكمية" and t0.cell(0, 1).text == "البند والكمية"  # ô gộp colspan=2
    assert t0.cell(1, 2).text == "1,250.00"
    paras = {p.text: p for p in d.paragraphs}
    ar = next(p for t, p in paras.items() if "العميل" in t)
    assert ar._p.pPr.find(qn("w:bidi")) is not None
    en = paras["Page two in English."]
    assert en._p.pPr is None or en._p.pPr.find(qn("w:bidi")) is None
    assert any(r.bold for p in d.paragraphs for r in p.runs if "رقم الفاتورة" in r.text)
    xml = d.element.xml
    assert xml.count('w:type="page"') == 1  # 2 trang nguồn -> 1 ngắt trang


def test_parse_pages():
    assert parse_pages("1-3,5", 10) == [1, 2, 3, 5]
    assert parse_pages("", 3) == [1, 2, 3]
    assert parse_pages("2-99", 4) == [2, 3, 4]


@pytest.fixture
def converter(tmp_path):
    (tmp_path / "m.jsonl").write_text("")
    (tmp_path / "c.yaml").write_text(yaml.safe_dump({"dataset": "m.jsonl", "models": [
        {"name": "fake", "adapter": "tests.test_app:FakeOCR"}]}))
    return Converter(tmp_path / "c.yaml", "fake")


def _scan_pdf(path, n=2):
    pages = []
    for i in range(n):
        im = Image.new("RGB", (600, 800), "white")
        ImageDraw.Draw(im).text((50, 50), f"scanned page {i + 1}", fill="black")
        pages.append(im)
    pages[0].save(path, save_all=True, append_images=pages[1:])


def test_scanned_pdf_goes_through_ocr(converter, tmp_path):
    _scan_pdf(tmp_path / "scan.pdf")
    res = converter.convert_file(tmp_path / "scan.pdf", tmp_path / "out", doc_type="table")
    assert [p.source for p in res.pages] == ["OCR", "OCR"]
    assert "[table_html]" in res.pages[0].text  # loại tài liệu "Có bảng" -> prompt cho bảng
    d = Document(res.docx)
    assert len(d.tables) == 4 and "Trang 1 — nguồn: OCR · model fake" in d.paragraphs[0].text


def test_image_input_and_page_range(converter, tmp_path):
    Image.new("RGB", (300, 300), "white").save(tmp_path / "a.png")
    assert [p.source for p in converter.convert_file(tmp_path / "a.png", tmp_path / "o").pages] == ["OCR"]
    _scan_pdf(tmp_path / "s.pdf", n=3)
    assert [p.index for p in converter.convert_file(tmp_path / "s.pdf", tmp_path / "o", pages="2-3").pages] == [2, 3]


def _digital_pdf(path, html):
    pw = pytest.importorskip("playwright.sync_api")
    try:
        with pw.sync_playwright() as p:
            try:
                b = p.chromium.launch()
            except Exception:
                b = p.chromium.launch(channel="chrome")
            pg = b.new_page()
            pg.set_content(html)
            pg.pdf(path=str(path), format="A4")
            b.close()
    except Exception as e:
        pytest.skip(f"không có Chromium: {e}")


def test_text_layer_used_for_english_but_not_arabic(converter, tmp_path):
    _digital_pdf(tmp_path / "en.pdf", "<p>" + "This contract is made between the parties below. " * 5 + "</p>")
    res = converter.convert_file(tmp_path / "en.pdf", tmp_path / "o")
    assert res.pages[0].source == "lớp chữ PDF" and "This contract is made" in res.pages[0].text
    # lớp chữ tiếng Ả Rập trong PDF hay bị đảo thứ tự -> luôn OCR
    _digital_pdf(tmp_path / "ar.pdf", "<p dir=rtl>" + "إنه في يوم الخميس تم الاتفاق بين الطرفين على ما يلي. " * 4 + "</p>")
    assert converter.convert_file(tmp_path / "ar.pdf", tmp_path / "o").pages[0].source == "OCR"
    # bật force_ocr: trang tiếng Anh cũng qua OCR
    assert converter.convert_file(tmp_path / "en.pdf", tmp_path / "o", force_ocr=True).pages[0].source == "OCR"


def test_gradio_app_builds(converter):
    pytest.importorskip("gradio")
    from ocrbench.app import build_app

    demo = build_app(converter)
    assert demo is not None


def test_parse_variant_and_merge():
    from ocrbench.convert import merge_params
    from ocrbench.speedtest import parse_variant

    p = parse_variant("max_new_tokens=4096,batch_size=4,stop_on_loop=true,max_pixels=1600000")
    assert p == {"max_new_tokens": 4096, "batch_size": 4, "stop_on_loop": True,
                 "processor_kwargs": {"max_pixels": 1600000}}
    merged = merge_params({"prompt": "x", "processor_kwargs": {"min_pixels": 1}}, p)
    assert merged["processor_kwargs"] == {"min_pixels": 1, "max_pixels": 1600000} and merged["prompt"] == "x"


class FlakyBatch(Adapter):
    """Batch nhiều trang thì lỗi (giả lập hết VRAM), từng trang thì được."""

    def predict(self, image, item):
        return Prediction(f"ok {item.id}", extra={"stopped_loop": item.id.endswith("p2")})

    def predict_batch(self, images, items):
        if len(images) > 1:
            raise RuntimeError("CUDA out of memory (giả lập)")
        return [self.predict(images[0], items[0])]


def test_batch_failure_falls_back_per_page(tmp_path):
    (tmp_path / "m.jsonl").write_text("")
    (tmp_path / "c.yaml").write_text(yaml.safe_dump({"dataset": "m.jsonl", "models": [
        {"name": "flaky", "adapter": "tests.test_app:FlakyBatch", "params": {"batch_size": 3}}]}))
    conv = Converter(tmp_path / "c.yaml", "flaky")
    _scan_pdf(tmp_path / "s.pdf", n=4)
    res = conv.convert_file(tmp_path / "s.pdf", tmp_path / "o")
    assert [p.index for p in res.pages] == [1, 2, 3, 4] and all(p.source == "OCR" for p in res.pages)
    assert conv.fallbacks >= 1
    assert res.pages[1].note == "model bị lặp, đã dừng sớm"
    assert "CẦN SOÁT" in Document(res.docx).element.xml
